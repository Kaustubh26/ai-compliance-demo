# orchestrator.py  (LOCAL ONEDRIVE VERSION WITH DOCX REPORT)

import os
from pathlib import Path
from typing import Dict, List

from dotenv import load_dotenv
from docx import Document
from pypdf import PdfReader

from agents.compliance_agent import run_compliance_analysis
from agents.resolver_agent import resolve_cross_document_issues

# Always load .env from project root
ENV_PATH = Path(__file__).resolve().parent / ".env"
load_dotenv(ENV_PATH)

# Local OneDrive-synced folders (set in .env)
POLICIES_DIR = os.getenv("LOCAL_POLICIES_DIR")
STANDARDS_DIR = os.getenv("LOCAL_STANDARDS_DIR")
REPORTS_DIR = os.getenv("LOCAL_REPORTS_DIR")


# =========================
# Helpers: file -> text
# =========================

def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_docx(path: Path) -> str:
    from docx import Document as DocxDocument  # avoid name clash
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: List[str] = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        parts.append(txt)
    return "\n".join(parts)


def load_file_as_text(path: Path) -> str:
    lower = path.name.lower()
    if lower.endswith(".txt"):
        return read_txt(path)
    if lower.endswith(".docx"):
        return read_docx(path)
    if lower.endswith(".pdf"):
        return read_pdf(path)
    raise ValueError(f"Unsupported file type for: {path.name} (only .txt, .docx, .pdf)")


def list_supported_files(folder: Path) -> List[Path]:
    if not folder.exists():
        raise FileNotFoundError(f"Folder does not exist: {folder}")
    if not folder.is_dir():
        raise NotADirectoryError(f"Not a directory: {folder}")

    exts = {".txt", ".docx", ".pdf"}
    return sorted([p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in exts])


# =========================
# Docx report helper
# =========================

def write_report_docx(report_text: str, report_path: Path) -> None:
    """
    Take the large LLM-generated report text and write it as a nicely
    formatted Word document with headings based on the === markers.
    """
    doc = Document()

    # Main title
    doc.add_heading("Global Compliance Report", level=1)

    current_section = None

    for raw_line in report_text.splitlines():
        line = raw_line.strip()
        if not line:
            # Blank line -> just a spacer
            doc.add_paragraph("")
            continue

        # Treat lines starting with === as section headings
        if line.startswith("===") and line.endswith("==="):
            # Example: "=== GLOBAL COMPLIANCE SUMMARY ==="
            heading_text = line.strip("=").strip()
            doc.add_heading(heading_text, level=2)
            current_section = heading_text
        else:
            # Normal paragraph or bullet line
            if line.startswith("- ") or line.startswith("• "):
                # Bullet item
                doc.add_paragraph(line.lstrip("-• ").strip(), style="List Bullet")
            else:
                # Normal paragraph
                doc.add_paragraph(line)

    # Ensure parent folder exists
    report_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(report_path))


# =========================
# Orchestrator
# =========================

def orchestrator() -> None:
    """
    Local-only pipeline using OneDrive-synced folders:

    - Read all standards from LOCAL_STANDARDS_DIR
    - Read all policies from LOCAL_POLICIES_DIR
    - Compare each policy against combined standards using Azure OpenAI
    - Use resolver to aggregate issues
    - Write global_compliance_report.docx into LOCAL_REPORTS_DIR
    """

    print("Using LOCAL_POLICIES_DIR:", POLICIES_DIR)
    print("Using LOCAL_STANDARDS_DIR:", STANDARDS_DIR)
    print("Using LOCAL_REPORTS_DIR:", REPORTS_DIR)

    if not POLICIES_DIR or not STANDARDS_DIR or not REPORTS_DIR:
        raise RuntimeError(
            "LOCAL_POLICIES_DIR, LOCAL_STANDARDS_DIR, LOCAL_REPORTS_DIR "
            "must be set in .env"
        )

    policies_folder = Path(POLICIES_DIR).expanduser()
    standards_folder = Path(STANDARDS_DIR).expanduser()
    reports_folder = Path(REPORTS_DIR).expanduser()
    reports_folder.mkdir(parents=True, exist_ok=True)

    print("\n🚀 Starting Multi-Agent Compliance System (Local OneDrive sync)")
    print("================================================================\n")

    # ------------------------------------------------------------------
    # 1. Load standards
    # ------------------------------------------------------------------
    print("📥 Loading standards from local OneDrive folder...")
    standards_files = list_supported_files(standards_folder)

    if not standards_files:
        print(f"⚠️ No standards files found in: {standards_folder}")
        combined_standards_text = ""
    else:
        standard_blocks: List[str] = []
        for std_path in standards_files:
            print(f"   - Reading standard: {std_path.name}")
            try:
                std_text = load_file_as_text(std_path)
            except ValueError as e:
                print(f"⚠️ Skipping standard '{std_path.name}': {e}")
                continue
            standard_blocks.append(
                f"==== STANDARD DOCUMENT: {std_path.name} ====\n{std_text}"
            )
        combined_standards_text = "\n\n".join(standard_blocks)

    # ------------------------------------------------------------------
    # 2. Load policies
    # ------------------------------------------------------------------
    print("\n📥 Loading policies from local OneDrive folder...")
    policy_files = list_supported_files(policies_folder)

    if not policy_files:
        print(f"⚠️ No policy files found in: {policies_folder}")
        return

    policy_data: Dict[str, dict] = {}

    for policy_path in policy_files:
        policy_name = policy_path.name
        print(f"\n📄 Running compliance analysis for: {policy_name}")

        try:
            policy_text = load_file_as_text(policy_path)
        except ValueError as e:
            print(f"⚠️ Skipping policy '{policy_name}': {e}")
            continue

        analysis = run_compliance_analysis(
            policy_text=policy_text,
            policy_name=policy_name,
            standards_text=combined_standards_text,
        )

        policy_data[policy_name] = {
            "policy_text": policy_text,
            "analysis": analysis,
        }

    if not policy_data:
        print("\n⚠️ No policies were successfully processed. Nothing to report.")
        return

    # ------------------------------------------------------------------
    # 3. Resolve cross-policy issues -> get final report text
    # ------------------------------------------------------------------
    print("\n🧠 Resolving cross-document issues...")
    resolved_report_text = resolve_cross_document_issues(policy_data, combined_standards_text)

    # ------------------------------------------------------------------
    # 4. Write final report as a Word document into OneDrive Reports folder
    # ------------------------------------------------------------------
    docx_report_path = reports_folder / "global_compliance_report.docx"
    print(f"\n📝 Writing final report to Word document: {docx_report_path}")
    write_report_docx(resolved_report_text, docx_report_path)

    print("\n✅ Compliance report written as a .docx file.")
    print("   OneDrive will sync it automatically.")
    print("   Report file:", docx_report_path)


if __name__ == "__main__":
    orchestrator()
